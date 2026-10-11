
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


workbook = Workbook()
sheet = workbook.active
sheet.title = "Semester_Results"

headers = [
    "Roll No",
    "Student Name",
    "Lab Marks (out of 50)",
    "Theory Marks (out of 50)",
    "Total Marks (out of 100)",
    "Result"
]

sheet.append(headers)

students = [
    [1, "Aarav", 42, 45],
    [2, "Diya", 35, 38],
    [3, "Rohan", 20, 15],
    [4, "Ananya", 46, 44],
    [5, "Kabir", 28, 25]
]

for student in students:
    roll_no, name, lab_marks, theory_marks = student
    total_marks = lab_marks + theory_marks
    result = "Pass" if total_marks >= 40 else "Fail"

    sheet.append([
        roll_no,
        name,
        lab_marks,
        theory_marks,
        total_marks,
        result
    ])

header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9E1F2"
)

for cell in sheet[1]:
    cell.font = Font(bold=True)
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

average_row = sheet.max_row + 1
sheet.cell(row=average_row, column=2, value="Class Average")

for column in range(3, 6):
    letter = sheet.cell(row=1, column=column).column_letter
    sheet.cell(
        row=average_row,
        column=column,
        value=f"=AVERAGE({letter}2:{letter}{average_row - 1})"
    )

sheet.cell(row=average_row, column=2).font = Font(bold=True)

for column in range(1, 7):
    sheet.column_dimensions[
        sheet.cell(row=1, column=column).column_letter
    ].width = 24

sheet.freeze_panes = "A2"

workbook.save("AIML_Batch_Grades.xlsx")

print("Student records processed successfully!")
print("Excel file saved as AIML_Batch_Grades.xlsx")